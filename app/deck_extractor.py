from __future__ import annotations

from io import BytesIO
import re
import zipfile

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from dataclasses import dataclass
import hashlib

MAX_UNCOMPRESSED_PPTX_BYTES = 100 * 1024 * 1024


class DeckExtractionError(Exception):
    """Raised when a PPTX cannot be safely or meaningfully extracted."""


def normalize_text(text: str) -> str:
    """
    Normalize whitespace without changing the semantic content.
    """
    if not text:
        return ""

    text = text.replace("\r\n", "\n")
    text = text.replace("\r", "\n")

    # collapse repeated spaces/tabs, but preserve line boundaries
    text = re.sub(r"[ \t]+", " ", text)

    lines = [line.strip() for line in text.split("\n")]

    # remove excessive blank lines
    cleaned_lines = []
    previous_blank = False

    for line in lines:
        is_blank = not line

        if is_blank and previous_blank:
            continue

        cleaned_lines.append(line)
        previous_blank = is_blank

    return "\n".join(cleaned_lines).strip()


def validate_pptx_container(data: bytes) -> None:
    """
    PPTX is a ZIP container.
    Perform a few cheap safety/integrity checks before python-pptx opens it.
    """
    if not data:
        raise DeckExtractionError("PPTX file is empty.")

    bio = BytesIO(data)

    if not zipfile.is_zipfile(bio):
        raise DeckExtractionError(
            "Attachment is not a valid PPTX/ZIP container."
        )

    bio.seek(0)

    with zipfile.ZipFile(bio) as archive:
        names = set(archive.namelist())

        if "ppt/presentation.xml" not in names:
            raise DeckExtractionError(
                "ZIP file does not contain a PowerPoint presentation."
            )

        total_uncompressed = sum(
            item.file_size for item in archive.infolist()
        )

        if total_uncompressed > MAX_UNCOMPRESSED_PPTX_BYTES:
            raise DeckExtractionError(
                "PPTX expands beyond the allowed size."
            )


def text_from_text_frame(shape) -> str:
    parts = []

    for paragraph in shape.text_frame.paragraphs:
        text = normalize_text(paragraph.text)

        if text:
            parts.append(text)

    return "\n".join(parts)


def text_from_table(shape) -> str:
    """
    Render a PowerPoint table as readable plain text.
    Each table row becomes one line.
    """
    rows = []

    for row in shape.table.rows:
        cells = [
            normalize_text(cell.text)
            for cell in row.cells
        ]

        if any(cells):
            rows.append(" | ".join(cells))

    return "\n".join(rows)


def collect_shape_blocks(shape, parent_top=0, parent_left=0):
    """
    Return:
        [
            (top, left, "some text"),
            ...
        ]

    Coordinates are used only to create a reasonable reading order.
    """
    top = parent_top + int(getattr(shape, "top", 0) or 0)
    left = parent_left + int(getattr(shape, "left", 0) or 0)

    blocks = []

    # Grouped PowerPoint objects
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for child in shape.shapes:
            blocks.extend(
                collect_shape_blocks(
                    child,
                    parent_top=top,
                    parent_left=left,
                )
            )

        return blocks

    # Tables
    if getattr(shape, "has_table", False):
        text = text_from_table(shape)

        if text:
            blocks.append((top, left, text))

        return blocks

    # Regular text boxes/placeholders
    if getattr(shape, "has_text_frame", False):
        text = text_from_text_frame(shape)

        if text:
            blocks.append((top, left, text))

    # Chart title, when exposed as text by python-pptx
    if getattr(shape, "has_chart", False):
        chart = shape.chart

        if chart.has_title:
            title = normalize_text(
                chart.chart_title.text_frame.text
            )

            if title:
                blocks.append((top, left, title))

    return blocks


def extract_pptx_text(data: bytes) -> dict:
    """
    Extract visible textual content slide-by-slide.

    Returns a dictionary containing:
        slide_count
        nonempty_slide_count
        char_count
        slides
        full_text
        warnings
    """
    validate_pptx_container(data)

    try:
        presentation = Presentation(BytesIO(data))
    except Exception as exc:
        raise DeckExtractionError(
            f"PowerPoint file could not be opened: {exc}"
        ) from exc

    slides = []
    warnings = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1,
    ):
        blocks = []

        for shape in slide.shapes:
            blocks.extend(collect_shape_blocks(shape))

        # approximate natural reading order:
        # top-to-bottom, then left-to-right
        blocks.sort(key=lambda item: (item[0], item[1]))

        texts = []
        seen = set()

        for _, _, raw_text in blocks:
            text = normalize_text(raw_text)

            if not text:
                continue

            # Avoid obvious duplicate blocks
            if text in seen:
                continue

            seen.add(text)
            texts.append(text)

        slide_text = "\n".join(texts).strip()

        if not slide_text:
            warnings.append(
                f"Slide {slide_number} contains no extractable text."
            )

        slides.append(
            {
                "slide_number": slide_number,
                "text": slide_text,
            }
        )

    nonempty_slide_count = sum(
        1 for slide in slides if slide["text"]
    )

    full_text_parts = []

    for slide in slides:
        full_text_parts.append(
            f"--- SLIDE {slide['slide_number']} ---"
        )

        if slide["text"]:
            full_text_parts.append(slide["text"])
        else:
            full_text_parts.append("[NO EXTRACTABLE TEXT]")

        full_text_parts.append("")

    full_text = "\n".join(full_text_parts).strip()

    real_text_char_count = sum(
        len(slide["text"])
        for slide in slides
    )

    return {
        "slide_count": len(slides),
        "nonempty_slide_count": nonempty_slide_count,
        "char_count": real_text_char_count,
        "slides": slides,
        "full_text": full_text,
        "warnings": warnings,
    }



MIN_DECK_TEXT_CHARS = 200


@dataclass(frozen=True)
class DeckWarning:
    message: str


@dataclass(frozen=True)
class DeckPage:
    page_number: int
    text: str

    @property
    def slide_number(self) -> int:
        return self.page_number

    @property
    def number(self) -> int:
        return self.page_number


# @dataclass(frozen=True)
# class ParsedDeck:
#     pages: list[DeckPage]
#     page_count: int
#     sha256: str
#     warnings: list[DeckWarning]
#     full_text: str
#     char_count: int

#     @property
#     def text(self) -> str:
#         return self.full_text

@dataclass(frozen=True)
class ParsedDeck:
    pages: list[DeckPage]
    page_count: int
    sha256: str
    warnings: list[DeckWarning]
    full_text: str
    char_count: int

    @property
    def text(self) -> str:
        return self.full_text

    def to_dict(self) -> dict:
        return {
            "pages": [
                {
                    "page_number": page.page_number,
                    "text": page.text,
                }
                for page in self.pages
            ],
            "page_count": self.page_count,
            "sha256": self.sha256,
            "warnings": [
                {
                    "message": warning.message,
                }
                for warning in self.warnings
            ],
            "full_text": self.full_text,
            "char_count": self.char_count,
        }

def parse_pptx_document(
    data: bytes,
) -> ParsedDeck:
    result = extract_pptx_text(data)

    if result["char_count"] < MIN_DECK_TEXT_CHARS:
        raise DeckExtractionError(
            "Too little extractable text was found in the deck."
        )

    pages: list[DeckPage] = []

    for index, slide in enumerate(
        result["slides"],
        start=1,
    ):
        slide_number = slide.get(
            "slide_number",
            slide.get("number", index),
        )

        slide_text = slide.get(
            "text",
            "",
        )

        pages.append(
            DeckPage(
                page_number=int(slide_number),
                text=slide_text,
            )
        )

    warnings = [
        DeckWarning(message=str(item))
        for item in result["warnings"]
    ]

    return ParsedDeck(
        pages=pages,
        page_count=result["slide_count"],
        sha256=hashlib.sha256(data).hexdigest(),
        warnings=warnings,
        full_text=result["full_text"],
        char_count=result["char_count"],
    )